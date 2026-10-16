from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

OUT = r"D:\jc2001\reports\phase2_supervisor_briefing_bilingual.docx"

def set_font(run, name="Calibri", size=10.5, color=None, bold=False, italic=False):
    run.font.name = name
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), "Microsoft YaHei" if name == "Calibri" else name)
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    if color:
        run.font.color.rgb = RGBColor(*color)

def shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tcPr.append(shd)

def set_cell_text(cell, text, bold=False, color=None):
    cell.text = ""
    p = cell.paragraphs[0]
    r = p.add_run(text)
    set_font(r, size=9, bold=bold, color=color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


doc = Document()
sec = doc.sections[0]
sec.top_margin = Cm(1.7)
sec.bottom_margin = Cm(1.7)
sec.left_margin = Cm(1.9)
sec.right_margin = Cm(1.9)

normal = doc.styles["Normal"]
normal.font.name = "Calibri"
normal.font.size = Pt(10.5)
normal._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), "Microsoft YaHei")
for style_name, size, color in [("Title", 22, (31, 78, 121)), ("Heading 1", 15, (31, 78, 121)), ("Heading 2", 12, (55, 96, 146))]:
    st = doc.styles[style_name]
    st.font.name = "Calibri"
    st.font.size = Pt(size)
    st.font.color.rgb = RGBColor(*color)
    st._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), "Microsoft YaHei")

def heading(text, level=1):
    p = doc.add_paragraph(style=f"Heading {level}")
    r = p.add_run(text)
    set_font(r, size=15 if level == 1 else 12, color=(31,78,121) if level == 1 else (55,96,146), bold=True)
    return p

def para(text="", style=None, color=None, italic=False, bold=False, space_after=5):
    p = doc.add_paragraph(style=style)
    p.paragraph_format.space_after = Pt(space_after)
    r = p.add_run(text)
    set_font(r, size=10.5, color=color, italic=italic, bold=bold)
    return p

def pair(en, zh):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run("EN  ")
    set_font(r, size=10.5, color=(31,78,121), bold=True)
    r = p.add_run(en)
    set_font(r, size=10.5)
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(0.35)
    p.paragraph_format.space_after = Pt(7)
    r = p.add_run("中文  ")
    set_font(r, size=10.5, color=(100,100,100), bold=True)
    r = p.add_run(zh)
    set_font(r, size=10.5, color=(80,80,80))

def bullet(text, color=None):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run(text)
    set_font(r, size=10, color=color)
    return p

# Title
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("Phase 2 / Cycle 2\nSupervisor Progress Briefing")
set_font(r, size=22, color=(31,78,121), bold=True)
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("第二阶段 / 第二周期导师汇报双语演讲稿")
set_font(r, size=16, color=(80,80,80), bold=True)

meta = doc.add_table(rows=6, cols=2)
meta.alignment = WD_TABLE_ALIGNMENT.CENTER
meta.style = "Table Grid"
rows = [
    ("Course / 课程", "JC2001 Introduction to Software Engineering"),
    ("Project / 项目", "Smart Study Assistant System (SmartStudy Console)"),
    ("Group / 小组", "Group 5"),
    ("Supervisor / 导师", "Dr. Shahzad Mumtaz"),
    ("Presenter / 汇报人", "Yuxuan Wu (50106070)"),
    ("Period / 周期", "Phase 2 / Cycle 2: 28 September–15 October 2026"),
]
for i, (a,b) in enumerate(rows):
    set_cell_text(meta.cell(i,0), a, bold=True, color=(31,78,121)); shade(meta.cell(i,0), "EAF2F8")
    set_cell_text(meta.cell(i,1), b)

doc.add_paragraph()
para("使用说明：英文部分是建议实际口述内容；中文部分用于理解、记忆和临场提示。建议时长约 8 分钟，现场演示另加 2 分钟。", color=(90,90,90), italic=True)

heading("一、开场 Opening | 约 30 秒", 1)
pair(
    "Good afternoon, Dr. Mumtaz. Thank you for meeting with me. I am Yuxuan Wu, the project manager of Group 5. Today I would like to report our Phase Two progress on the Smart Study Assistant System, or SmartStudy Console. I will briefly cover our requirements work, system modelling, working proof of concept, quality evidence, and the next risks we need to manage.",
    "下午好，Mumtaz 博士。感谢您参加这次会议。我是第五组的项目经理吴宇轩。今天我想汇报 Smart Study Assistant System，也就是 SmartStudy Console，在第二阶段的进展。我会依次介绍需求工作、系统建模、可运行原型、质量证据，以及接下来需要管理的风险。"
)

heading("二、阶段定位 Phase Positioning | 约 45 秒", 1)
pair(
    "Phase One focused on team formation, problem definition and the project proposal. Phase Two was the transition from a written proposal to an engineered and testable system. Our central design principle was that the system should diagnose why a student selected a distractor, rather than simply reveal the correct answer.",
    "第一阶段主要完成了团队组建、问题定义和项目提案。第二阶段的重点，是把纸面提案转化为可运行、可测试的软件系统。我们的核心设计原则是：系统要诊断学生为什么会选择某个干扰项，而不是直接告诉学生正确答案。"
)
pair(
    "This addresses two weaknesses of existing tools: conventional quizzes often provide only right-or-wrong feedback, while unrestricted large language models may reveal answers too early or produce unsupported explanations. Our solution combines option-level misconception diagnosis, knowledge-graph-style grounding, and Socratic guidance.",
    "这针对现有工具的两个问题：传统测验通常只提供对错反馈，而没有约束的大语言模型可能过早泄露答案，或者给出缺乏依据的解释。我们的方案把选项级错误诊断、知识图谱式的概念关联和苏格拉底式引导结合起来。"
)

heading("三、第二阶段完成的四项主要工作 Four Main Deliverables | 约 4 分钟", 1)
heading("1. Requirements Engineering 需求工程", 2)
pair(
    "First, we established a formal requirements baseline. We decomposed the system into eight Cockburn-style use cases, from student login and quiz participation to distractor diagnosis, Socratic tutoring and competency tracking. We also created a MoSCoW traceability structure covering functional requirements FR-01 to FR-08 and non-functional requirements NFR-01 to NFR-06.",
    "第一，我们建立了正式的需求基线。我们按照 Cockburn 用例方法拆分出八个用例，覆盖学生登录、答题、干扰项诊断、苏格拉底辅导和能力追踪等流程。同时建立了 MoSCoW 需求追溯结构，覆盖 FR-01 到 FR-08 的功能需求以及 NFR-01 到 NFR-06 的非功能需求。"
)
pair(
    "The important point is that we did not treat every feature as a must-have. We explicitly prioritised the reliable single-turn diagnostic experience for this cycle, while multi-turn tutoring, adaptive variants and cloud graph persistence remain planned extensions.",
    "重要的是，我们没有把所有功能都标记为必须完成。本周期优先保证可靠的单轮诊断体验；多轮导师、变式题生成和云端图谱持久化则明确列为后续扩展。"
)

heading("2. UML and Architecture UML 与系统架构", 2)
pair(
    "Second, we produced a UML 2.5 architecture package: a system context model, a global use-case model, a core sequence diagram for the diagnosis flow, and a domain class model. These diagrams helped us make the system boundary, actors, services and data contracts explicit before further implementation.",
    "第二，我们完成了 UML 2.5 架构模型包，包括系统上下文模型、全局用例模型、诊断流程核心时序图，以及领域类模型。这些图帮助我们在继续开发前明确系统边界、参与者、服务和数据契约。"
)
pair(
    "In the Week Six modelling exercise, we also practised the semantics of include and extend relationships using a university library example. This was useful because it reinforced the difference between mandatory shared behaviour and conditional extension behaviour, which we then applied when reviewing our own diagrams.",
    "在第六周的建模练习中，我们还用大学图书馆系统练习了 include 和 extend 关系。这帮助我们进一步区分必需的公共行为与有条件触发的扩展行为，并将这个原则应用到项目自身的图中。"
)

heading("3. Working PoC MVP 可运行原型", 2)
pair(
    "Third, we delivered an end-to-end proof of concept. The backend uses a modular FastAPI structure with routers, services and Pydantic data models. The current live routes are health, quiz and diagnosis. The frontend is an interactive console with a dark visual language, a command palette, a comparison sandbox, and a five-dimension competency radar.",
    "第三，我们完成了端到端可运行的概念验证原型。后端采用模块化 FastAPI 结构，分为 routers、services 和 Pydantic 数据模型。目前可运行的接口包括 health、quiz 和 diagnosis。前端是交互式控制台，包含暗色视觉风格、命令面板、对比沙盒和五维能力雷达。"
)
pair(
    "The seed dataset contains twelve questions across four disciplines: law, computer science, economics and software engineering. Each discipline has three questions. The questions contain thirty-six option-level distractor traps. All twelve questions also have a small graph structure: sixty concept nodes and forty-eight directed edges in total.",
    "种子数据集包含十二道题，覆盖法律、计算机科学、经济学和软件工程四个学科，每个学科三道题。错误选项共对应三十六个认知陷阱。十二道题都带有轻量级图结构，总计六十个概念节点和四十八条有向边。"
)
pair(
    "For resilience, the diagnosis service has two modes. When the external Qwen service is available, it can use the live path with a four-second timeout. When credentials are missing or the network fails, it falls back to the curated local seed data instead of returning an unhandled server error. This makes the prototype demonstrable in an offline classroom environment.",
    "在可靠性方面，诊断服务有两种模式。外部 Qwen 服务可用时，系统走在线路径，并设置四秒超时；缺少凭据或网络失败时，系统回退到本地整理好的种子数据，而不是抛出未处理的服务器错误。这让原型在没有网络的课堂环境中也可以演示。"
)

heading("4. Verification and QA 验证与质量保证", 2)
pair(
    "Fourth, we built an automated quality gate. The current mainline collects and passes 114 pytest cases with zero failures. The suites cover adversarial backend inputs, quiz contracts, CORS behaviour, full-stack stress scenarios, diagnosis logic, health checks and frontend asset delivery.",
    "第四，我们建立了自动化质量门禁。目前主干可以收集并通过 114 个 pytest 用例，失败数为零。测试覆盖后端恶意输入、测验契约、CORS 行为、全栈压力场景、诊断逻辑、健康检查以及前端静态资源交付。"
)
pair(
    "The security-oriented cases include SQL injection strings, XSS payloads, path traversal probes, malformed JSON and extreme input lengths. We also tested concurrent diagnosis and quiz requests. The key result for this phase is not only that the happy path works, but that common failure paths are handled without an unhandled HTTP 500 response.",
    "安全相关用例包括 SQL 注入字符串、XSS 载荷、路径穿越、错误 JSON 和超长输入。我们还测试了并发诊断和并发答题请求。本阶段的关键成果不仅是正常路径能够运行，还包括常见失败路径不会导致未处理的 HTTP 500 错误。"
)

heading("四、工程治理与本周期收尾 Engineering Governance | 约 45 秒", 1)
pair(
    "We also improved the engineering process. Work was organised through feature branches and reviewable commits. Before integration, we ran the full test suite. This week, we landed the Cycle Two progress report and the Week Six UML lab artefacts, merged the visual overhaul into develop, and then merged develop into main locally. The final main branch still passes all 114 tests.",
    "我们也完善了工程流程。工作通过功能分支和可审查提交来组织。每次集成前都运行完整测试套件。本周我们提交了第二周期进度报告和第六周 UML 实验材料，将前端视觉改版合并到 develop，再把 develop 合并到 main。本地最终 main 分支仍然通过全部 114 个测试。"
)

# Metrics table
heading("五、关键数据 Key Evidence", 1)
t = doc.add_table(rows=1, cols=3)
t.style = "Table Grid"; t.alignment = WD_TABLE_ALIGNMENT.CENTER
for c, txt in enumerate(["Metric / 指标", "Verified result / 已核验结果", "Evidence / 证据"]):
    set_cell_text(t.cell(0,c), txt, bold=True, color=(255,255,255)); shade(t.cell(0,c), "1F4E79")
metrics = [
    ("Automated QA", "114 passed, 0 failures", "python -m pytest -q"),
    ("Seed questions", "12 questions; 4 subjects", "backend/app/data/seed_questions.py"),
    ("Distractor traps", "36 option-level traps", "seed question data"),
    ("Graph seed", "12 graphs; 60 nodes; 48 edges", "Question model graph field"),
    ("API surface", "health / quiz / diagnose", "backend/app/routers/"),
    ("UML package", "Context, Use Case, Sequence, Class", "reports/figures/"),
    ("Repository history", "47 commits across refs", "git log --all"),
]
for row in metrics:
    cells = t.add_row().cells
    for i, val in enumerate(row): set_cell_text(cells[i], val)

doc.add_paragraph()

heading("六、建议现场演示 Live Demo Cue | 约 2 分钟", 1)
pair(
    "For a short demonstration, I would first show the diagnosis interface and select a plausible wrong option. The system should identify the distractor trap, show the related prerequisite concept, and provide a guiding question without exposing the answer letter. I would then show the offline fallback by running the same flow without depending on an external API key.",
    "现场演示时，我会先打开诊断界面，选择一个具有迷惑性的错误选项。系统应该识别干扰项陷阱，显示相关先修概念，并给出引导性问题，而不是直接显示答案字母。然后我会在不依赖外部 API key 的情况下重复这个流程，展示离线回退。"
)
para("演示命令 / Demo commands", bold=True, color=(31,78,121))
for cmd in ["python -m pytest -q", "python backend/run.py", "打开 frontend/index.html（或使用后端打印的 http://127.0.0.1:8000 地址）"]:
    bullet(cmd)

heading("七、当前风险与需要导师确认的事项 Risks and Supervisor Questions", 1)
pair(
    "There are three points I want to state transparently. First, the local main branch is complete and tested, but the final push to GitHub is currently blocked by a local network and proxy problem. Second, some team artefacts exist on member fork branches and need path consolidation before they can be treated as part of the mainline evidence package. Third, the current graph is a validated in-memory seed graph; live Neo4j or AuraDB integration is a next-stage task, not something I want to overclaim for this cycle.",
    "有三点我想透明说明。第一，本地 main 分支已经完成并通过测试，但最终 push 到 GitHub 目前受到本地网络和代理问题阻断。第二，部分组员材料存在于 fork 分支中，在被视为主干证据包之前还需要统一路径并合并。第三，目前的图谱是经过验证的内存种子图；Neo4j 或 AuraDB 的真实集成属于下一阶段，我不想把它夸大成已经完成。"
)
pair(
    "My main question for you is: for the upcoming architecture baseline, would you prefer us to demonstrate a live Neo4j instance now, or is a validated ontology, deterministic local graph and clear integration plan sufficient for the current milestone?",
    "我想请您确认一个重点：对于即将到来的架构基线，您希望我们现在就演示真实运行的 Neo4j 实例，还是说在当前里程碑中，经过验证的本体设计、确定性的本地图结构和清晰的集成计划已经足够？"
)

heading("八、下一步 Next Steps | 约 45 秒", 1)
pair(
    "Our immediate priorities are to complete the UML and Neo4j ontology baseline, consolidate the outstanding fork artefacts, and prepare Project Update One. After that, Phase Three will focus on live Neo4j or AuraDB integration, a multi-turn Socratic state machine, adaptive question variants, and broader user evaluation.",
    "我们近期的优先事项是完成 UML 与 Neo4j 本体基线、整合尚未进入主干的 fork 材料，并准备 Project Update One。之后第三阶段将重点完成真实 Neo4j 或 AuraDB 集成、多轮苏格拉底状态机、自适应变式题，以及更广泛的用户评估。"
)
pair(
    "Thank you. To summarise in one sentence: Phase Two transformed our proposal into a testable, resilient and demonstrable MVP, while also showing us exactly which architecture and integration tasks must be completed next.",
    "谢谢您。用一句话总结：第二阶段把项目提案转化成了可测试、可容错、可演示的 MVP，同时也明确了下一阶段必须完成的架构和集成任务。"
)

heading("九、导师可能追问的简答 Q&A Preparation", 1)
qa = [
    ("Why is this better than a normal quiz? / 为什么不只是普通测验？", "Because the system models the reason behind a wrong option and gives a Socratic prompt instead of only displaying correct or incorrect. / 因为系统建模的是错误选项背后的原因，并给出苏格拉底式引导，而不是只显示对错。"),
    ("Is the graph already Neo4j? / 现在已经是 Neo4j 了吗？", "Not yet in the mainline MVP. The current graph is a validated local graph structure. Neo4j or AuraDB integration is the next architecture milestone. / 主干 MVP 目前还不是。现在是经过验证的本地图结构，Neo4j 或 AuraDB 集成属于下一架构里程碑。"),
    ("How do you know it is robust? / 如何证明系统可靠？", "The current mainline collects 114 pytest cases and all pass, including adversarial, CORS, stress and fallback scenarios. / 当前主干收集 114 个 pytest 用例并全部通过，涵盖对抗性输入、CORS、压力和回退场景。"),
    ("What is the biggest current limitation? / 当前最大局限是什么？", "The MVP is still primarily single-turn and seed-data driven. Multi-turn state, cloud graph persistence and broader user evaluation remain open. / MVP 仍以单轮交互和种子数据为主，多轮状态、云端图谱持久化和更广泛的用户评估尚未完成。"),
    ("What will you deliver next? / 下一步交付什么？", "The UML and ontology baseline, consolidated repository evidence, Project Update One, and then the live graph integration plan. / UML 与本体基线、整合后的仓库证据、Project Update One，以及后续真实图谱集成计划。"),
]
for q, a in qa:
    para(q, bold=True, color=(31,78,121), space_after=2)
    para(a, color=(70,70,70), space_after=7)

heading("十、汇报时的三条底线 Three Accuracy Rules", 1)
for item in [
    "可以明确说：114 个测试在本地主干全部通过；12 道题、4 个学科、36 个陷阱、60 个节点和 48 条边均已从代码核验。",
    "不要说：Neo4j/AuraDB 已经完成；当前是本地种子图，真实云图集成属于下一阶段。",
    "不要回避：GitHub 最终 push 受网络阻断；如果导师现场查看远程仓库，应说明本地 main 已完成但远程同步待恢复网络后执行。",
]: bullet(item)

heading("附录：证据文件 Evidence References", 1)
refs = [
    ("Phase 2 progress report", "reports/phase2_cycle2_progress_report.md"),
    ("Backend routes and models", "backend/app/routers/ and backend/app/models/"),
    ("Seed questions and graph data", "backend/app/data/seed_questions.py"),
    ("Automated tests", "tests/"),
    ("UML figures", "reports/figures/"),
    ("Week 6 modelling guide", "reports/week6_lab_submission_guide.md"),
    ("Runtime entry point", "backend/run.py"),
]
t = doc.add_table(rows=1, cols=2); t.style="Table Grid"; t.alignment=WD_TABLE_ALIGNMENT.CENTER
for i, txt in enumerate(["Evidence / 证据", "Path / 路径"]):
    set_cell_text(t.cell(0,i), txt, bold=True, color=(255,255,255)); shade(t.cell(0,i), "1F4E79")
for a,b in refs:
    cells=t.add_row().cells; set_cell_text(cells[0],a); set_cell_text(cells[1],b)

# Footer
for section in doc.sections:
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = footer.add_run("JC2001 Group 5 | Phase 2 / Cycle 2 Supervisor Briefing")
    set_font(r, size=8, color=(120,120,120))

doc.save(OUT)
print(OUT)
